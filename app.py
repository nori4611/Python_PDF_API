import io
import json
import os
import secrets
import subprocess
import tempfile
from pathlib import Path

from flask import Flask, jsonify, request, send_file
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from werkzeug.utils import secure_filename


app = Flask(__name__)


@app.get("/")
def home():
    return jsonify(
        {
            "status": "online",
            "service": "Generic PowerPoint to PDF API",
            "endpoints": {
                "health": "/health",
                "generate": "/generate-document",
            },
        }
    )


@app.get("/health")
def health():
    return jsonify(
        {
            "status": "ok",
            "service": "pptx-pdf-worker",
        }
    )


def is_authorized():
    expected_key = os.getenv("API_KEY", "").strip()
    received_key = request.headers.get("X-API-Key", "").strip()

    return bool(
        expected_key
        and received_key
        and secrets.compare_digest(expected_key, received_key)
    )


def normalize_value(value):
    if value is None:
        return ""

    if isinstance(value, bool):
        return "true" if value else "false"

    return str(value).strip()


def read_record_data():
    """Read header1-header10 from record_json or form-data fields."""
    raw_record_json = request.form.get("record_json", "").strip()
    supplied_data = {}

    if raw_record_json:
        try:
            supplied_data = json.loads(raw_record_json)
        except json.JSONDecodeError as error:
            raise ValueError(f"record_json is not valid JSON: {error.msg}") from error

        if not isinstance(supplied_data, dict):
            raise ValueError("record_json must contain a JSON object")

    # Make JSON keys case-insensitive, for example Header1 or header1.
    normalized_input = {
        str(key).strip().lower(): value
        for key, value in supplied_data.items()
    }

    record_data = {}

    for number in range(1, 11):
        key = f"header{number}"

        value = normalized_input.get(key)

        # Also accept separate form-data fields as a fallback.
        if value is None:
            value = request.form.get(key)

        record_data[key] = normalize_value(value)

    # Temporary backward compatibility with the previous certificate workflow.
    if not record_data["header1"]:
        record_data["header1"] = request.form.get("name", "").strip()

    if not record_data["header2"]:
        record_data["header2"] = request.form.get("ic_number", "").strip()

    if not record_data["header3"]:
        record_data["header3"] = request.form.get("email", "").strip()

    return record_data


def build_replacements(record_data):
    replacements = {}

    for key, value in record_data.items():
        number = key.removeprefix("header")

        replacements[f"<header{number}>"] = value
        replacements[f"<Header{number}>"] = value
        replacements[f"{{{{header{number}}}}}"] = value
        replacements[f"{{{{Header{number}}}}}"] = value

    # Keep old certificate templates working during the migration.
    replacements.update(
        {
            "<name>": record_data["header1"],
            "<ic_number>": record_data["header2"],
            "<email>": record_data["header3"],
            "{{name}}": record_data["header1"],
            "{{ic_number}}": record_data["header2"],
            "{{email}}": record_data["header3"],
        }
    )

    return replacements


def replace_paragraph_text(paragraph, replacements):
    """Replace placeholders even when PowerPoint splits them across runs."""
    original_text = "".join(run.text for run in paragraph.runs)

    if not original_text:
        return

    updated_text = original_text

    for placeholder, value in replacements.items():
        updated_text = updated_text.replace(placeholder, value)

    if updated_text == original_text:
        return

    if paragraph.runs:
        # Preserve the formatting of the first run.
        paragraph.runs[0].text = updated_text

        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = updated_text


def replace_shape_text(shape, replacements):
    """Replace placeholders in text boxes, grouped shapes and table cells."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child_shape in shape.shapes:
            replace_shape_text(child_shape, replacements)

    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            replace_paragraph_text(paragraph, replacements)

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    replace_paragraph_text(paragraph, replacements)


def convert_pptx_to_pdf(input_pptx, output_directory, profile_directory):
    profile_directory.mkdir(parents=True, exist_ok=True)
    profile_uri = profile_directory.resolve().as_uri()

    result = subprocess.run(
        [
            "libreoffice",
            f"-env:UserInstallation={profile_uri}",
            "--headless",
            "--nologo",
            "--nodefault",
            "--nofirststartwizard",
            "--convert-to",
            "pdf:impress_pdf_Export",
            "--outdir",
            str(output_directory),
            str(input_pptx),
        ],
        check=True,
        capture_output=True,
        text=True,
        timeout=120,
        env={
            **os.environ,
            "HOME": str(output_directory),
        },
    )

    return result


# New generic endpoint.
# The old endpoint is retained temporarily for backward compatibility.
@app.post("/generate-document")
@app.post("/generate-certificate")
def generate_document():
    if not is_authorized():
        return jsonify(
            {
                "success": False,
                "message": "Unauthorized",
            }
        ), 401

    if "pptx_file" not in request.files:
        return jsonify(
            {
                "success": False,
                "message": "pptx_file is required",
            }
        ), 400

    execution_id = request.form.get("execution_id", "").strip()

    if not execution_id:
        return jsonify(
            {
                "success": False,
                "message": "execution_id is required",
            }
        ), 400

    safe_execution_id = secure_filename(execution_id)

    if not safe_execution_id:
        return jsonify(
            {
                "success": False,
                "message": "execution_id is invalid",
            }
        ), 400

    uploaded_file = request.files["pptx_file"]
    uploaded_name = uploaded_file.filename or ""

    if not uploaded_name.lower().endswith(".pptx"):
        return jsonify(
            {
                "success": False,
                "message": "pptx_file must be a .pptx file",
            }
        ), 400

    try:
        record_data = read_record_data()
    except ValueError as error:
        return jsonify(
            {
                "success": False,
                "execution_id": execution_id,
                "message": str(error),
            }
        ), 400

    header_count = sum(
        1 for value in record_data.values() if value != ""
    )

    if header_count == 0:
        return jsonify(
            {
                "success": False,
                "execution_id": execution_id,
                "message": "At least one header value is required",
            }
        ), 400

    replacements = build_replacements(record_data)

    try:
        with tempfile.TemporaryDirectory() as temporary_directory:
            workdir = Path(temporary_directory)

            input_path = workdir / "template.pptx"
            output_pptx = workdir / f"{safe_execution_id}.pptx"
            output_pdf = workdir / f"{safe_execution_id}.pdf"
            profile_directory = workdir / "libreoffice_profile"

            uploaded_file.save(input_path)

            presentation = Presentation(input_path)

            for slide in presentation.slides:
                for shape in slide.shapes:
                    replace_shape_text(shape, replacements)

            presentation.save(output_pptx)

            conversion = convert_pptx_to_pdf(
                output_pptx,
                workdir,
                profile_directory,
            )

            if not output_pdf.exists():
                converter_message = (
                    conversion.stderr.strip()
                    or conversion.stdout.strip()
                    or "No output was returned"
                )

                raise RuntimeError(
                    "LibreOffice did not produce the PDF file. "
                    f"Converter output: {converter_message}"
                )

            # Read into memory before TemporaryDirectory is deleted.
            pdf_buffer = io.BytesIO(output_pdf.read_bytes())
            pdf_buffer.seek(0)

            response = send_file(
                pdf_buffer,
                mimetype="application/pdf",
                as_attachment=True,
                download_name=f"{safe_execution_id}.pdf",
            )

            response.headers["X-Execution-ID"] = execution_id
            response.headers["X-Header-Count"] = str(header_count)

            return response

    except subprocess.TimeoutExpired:
        return jsonify(
            {
                "success": False,
                "execution_id": execution_id,
                "message": "PDF conversion timed out",
            }
        ), 504

    except subprocess.CalledProcessError as error:
        app.logger.exception("LibreOffice conversion failed")

        converter_message = (
            (error.stderr or "").strip()
            or (error.stdout or "").strip()
            or "LibreOffice conversion failed"
        )

        return jsonify(
            {
                "success": False,
                "execution_id": execution_id,
                "message": converter_message,
            }
        ), 500

    except Exception as error:
        app.logger.exception("PDF generation failed")

        return jsonify(
            {
                "success": False,
                "execution_id": execution_id,
                "message": str(error),
            }
        ), 500


if __name__ == "__main__":
    port = int(os.getenv("PORT", "10000"))
    app.run(host="0.0.0.0", port=port)
